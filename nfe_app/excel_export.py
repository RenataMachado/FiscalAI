"""
Geração de planilhas Excel: levantamento de notas e cálculo de pagamento com desconto.
"""
import io

def gerar_excel_levantamento(lista_dados):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Balanço de Notas"

    azul_cabecalho = PatternFill(start_color="005B96", end_color="005B96", fill_type="solid")
    fundo_cinza = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
    fonte_branca = Font(color="FFFFFF", bold=True)
    fonte_normal = Font(color="333333")
    fonte_bold = Font(color="333333", bold=True)
    
    borda_cinza = Border(
        left=Side(style='thin', color="D3D3D3"),
        right=Side(style='thin', color="D3D3D3"),
        top=Side(style='thin', color="D3D3D3"),
        bottom=Side(style='thin', color="D3D3D3")
    )

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center', wrap_text=True)

    cabecalho = [
        "Empresa", "NFE", "Contrato", "Contrato SP Águas", "Processo SEI", 
        "Emissão", "Vencimento", "CNPJ Emitente", 
        "Bruto (R$)", "ISS (R$)", "IRF (R$)", "Líquido (R$)"
    ]
    
    ws.append(cabecalho)
    for col_num in range(1, len(cabecalho) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = azul_cabecalho
        cell.font = fonte_branca
        cell.alignment = align_center
        cell.border = borda_cinza

    tot_bruto = tot_iss = tot_irf = tot_liq = 0.0

    for row_idx, item in enumerate(lista_dados, start=2):
        bruto = float(item.get("valor_total") or 0.0)
        iss = float(item.get("valor_iss") or 0.0)
        ir = float(item.get("valor_ir") or 0.0)
        liq = float(item.get("valor_liquido") or 0.0)

        tot_bruto += bruto
        tot_iss += iss
        tot_irf += ir
        tot_liq += liq

        linha = [
            str(item.get("Arquivo", "")),
            str(item.get("numero_nfe", "")),
            str(item.get("numero_contrato", "")),
            str(item.get("numero_contrato_spaguas", "")),
            str(item.get("processo_sei", "")),
            str(item.get("data_emissao", "")),
            str(item.get("vencimento", "")),
            str(item.get("cnpj_emitente", "")),
            bruto, iss, ir, liq
        ]
        
        ws.append(linha)
        
        for col_num in range(1, len(linha) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.font = fonte_normal
            cell.border = borda_cinza
            
            if row_idx % 2 == 1: 
                cell.fill = fundo_cinza
            
            if col_num == 1: 
                cell.alignment = align_left
            elif col_num >= 9: 
                cell.alignment = align_right
                cell.number_format = '#,##0.00'
            else: 
                cell.alignment = align_center
    
    row_tot = ws.max_row + 1
    ws.cell(row=row_tot, column=1, value="TOTAL")
    ws.cell(row=row_tot, column=9, value=tot_bruto)
    ws.cell(row=row_tot, column=10, value=tot_iss)
    ws.cell(row=row_tot, column=11, value=tot_irf)
    ws.cell(row=row_tot, column=12, value=tot_liq)

    for col_num in [1, 9, 10, 11, 12]:
        cell = ws.cell(row=row_tot, column=col_num)
        cell.font = fonte_bold
        cell.fill = PatternFill(start_color="E6EDF5", end_color="E6EDF5", fill_type="solid")
        cell.border = borda_cinza
        if col_num >= 9:
            cell.alignment = align_right
            cell.number_format = '#,##0.00'
        else:
            cell.alignment = align_center

    larguras = {
        'A': 40, 'B': 15, 'C': 20, 'D': 20, 'E': 25, 
        'F': 15, 'G': 15, 'H': 20, 'I': 18, 'J': 18, 
        'K': 18, 'L': 18
    }
    for col_letra, largura in larguras.items():
        ws.column_dimensions[col_letra].width = largura

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def gerar_excel_calculo_pagto(lista_dados, processo_sei_filtro=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Pagamento"

    # Filtro opcional por Processo SEI
    if processo_sei_filtro and processo_sei_filtro != "Todos":
        lista_dados = [item for item in lista_dados if str(item.get("processo_sei", "")).strip() == str(processo_sei_filtro).strip()]

    font_bold = Font(bold=True)
    font_red_bold = Font(color="FF0000", bold=True)
    align_center = Alignment(horizontal='center', vertical='center')
    border_thin = Border(
        left=Side(style='thin'), right=Side(style='thin'), 
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    if not lista_dados:
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    primeiro = lista_dados[0]
    empresa = str(primeiro.get("Arquivo", ""))
    
    tc_num = str(primeiro.get("numero_contrato_spaguas", ""))
    if not tc_num or tc_num == "-" or tc_num == "":
        tc_num = str(primeiro.get("numero_contrato", ""))
        
    cnpj = str(primeiro.get("cnpj_emitente", ""))
    natureza = str(primeiro.get("natureza_despesa", ""))
    ne = str(primeiro.get("numero_ne", ""))
    fonte = str(primeiro.get("fonte_recurso", ""))
    sei = str(primeiro.get("processo_sei", ""))
    ug_lida = str(primeiro.get("ug", ""))
    assunto_lido = str(primeiro.get("assunto", ""))

    ws.merge_cells('A1:G1')
    ws['A1'] = "CALCULO PARA PAGTO DE FATURAS C/DESCONTO"
    ws['A1'].font = font_bold
    ws['A1'].alignment = align_center

    ws['A2'] = "T.C. nº"
    ws.merge_cells('B2:C2')
    ws['B2'] = tc_num
    ws.merge_cells('D2:G2')
    ws['D2'] = f"ASSUNTO: {assunto_lido}" if assunto_lido and assunto_lido != "nan" else "ASSUNTO: PRESTAÇÃO DE SERVIÇOS TÉCNICOS ESPECIALIZADOS DE SUPERVISÃO, FISCALIZAÇÃO"
    ws['D2'].font = Font(size=9)
    
    ws['A3'] = "UG:"
    ws['B3'] = ug_lida if ug_lida and ug_lida != "nan" else "262103"
    ws['B3'].font = font_red_bold
    ws['B3'].alignment = align_center

    ws['A4'] = "EMPRESA"
    ws['A4'].alignment = align_center
    ws.merge_cells('B4:C4')
    ws['B4'] = empresa
    ws['D4'] = "NATUREZA"
    ws['D4'].alignment = align_center
    ws['E4'] = natureza
    ws['E4'].alignment = align_center
    ws['F4'] = "CNPJ:"
    ws['F4'].alignment = align_center
    ws['G4'] = cnpj
    ws['G4'].alignment = align_center

    ws['A5'] = "NE"
    ws['A5'].alignment = align_center
    ws['B5'] = ne
    ws['B5'].font = font_bold
    ws['B5'].alignment = align_center
    ws['C5'] = fonte
    ws['C5'].alignment = align_center
    ws['D5'] = "SEI"
    ws['D5'].alignment = align_center
    ws['E5'] = sei
    ws['E5'].alignment = align_center
    ws['F5'] = "MEDIÇÃO Nº"
    ws['F5'].alignment = align_center
    ws['G5'] = "" 

    ws['D6'] = "PROC"
    ws['D6'].alignment = align_center
    ws['E6'] = "20231599433" 

    for r in range(1, 7):
        for c in range(1, 8):
            ws.cell(row=r, column=c).border = border_thin

    cols = ["NFE", "BRUTO", "INSS", "ISS", "CAUÇÃO", "IRF", "LIQUIDO"]
    for i, col_name in enumerate(cols, 1):
        cell = ws.cell(row=7, column=i)
        cell.value = col_name
        cell.font = font_bold
        cell.alignment = align_center
        cell.border = border_thin

    row_idx = 8
    tot_bruto = tot_iss = tot_irf = tot_liq = 0.0

    for item in lista_dados:
        nfe = str(item.get("numero_nfe", ""))
        bruto = float(item.get("valor_total", 0.0))
        inss = 0.0
        iss = float(item.get("valor_iss", 0.0))
        caucao = 0.0
        irf = float(item.get("valor_ir", 0.0))
        liq = float(item.get("valor_liquido", 0.0))
        
        tot_bruto += bruto
        tot_iss += iss
        tot_irf += irf
        tot_liq += liq
        
        valores = [nfe, bruto, inss, iss, caucao, irf, liq]
        
        for i, val in enumerate(valores, 1):
            cell = ws.cell(row=row_idx, column=i)
            cell.border = border_thin
            if i == 1:
                cell.value = val
                cell.alignment = align_center
            else:
                if val == 0.0:
                    cell.value = "-"
                    cell.alignment = align_center
                else:
                    cell.value = val
                    cell.number_format = '#,##0.00'
        row_idx += 1

    for i, val in enumerate(["TOTAL", tot_bruto, "-", tot_iss, "-", tot_irf, tot_liq], 1):
        cell = ws.cell(row=row_idx, column=i)
        cell.value = val
        cell.font = font_bold
        cell.border = border_thin
        if isinstance(val, float): 
            cell.number_format = '#,##0.00'
        else:
            cell.alignment = align_center
    row_idx += 1

    for i, val in enumerate([fonte, tot_bruto, "-", tot_iss, "-", tot_irf, tot_liq], 1):
        cell = ws.cell(row=row_idx, column=i)
        cell.value = val
        cell.font = font_bold
        cell.border = border_thin
        if isinstance(val, float): 
            cell.number_format = '#,##0.00'
        else:
            cell.alignment = align_center
    row_idx += 1

    for i, val in enumerate(["", tot_bruto, "-", tot_iss, "-", tot_irf, tot_liq], 1):
        cell = ws.cell(row=row_idx, column=i)
        cell.value = val
        cell.font = font_bold
        cell.border = border_thin
        if isinstance(val, float): 
            cell.number_format = '#,##0.00'
        else:
            cell.alignment = align_center
    row_idx += 1

    ws.cell(row=row_idx, column=3).value = "NLRETINSS"
    ws.cell(row=row_idx, column=4).value = "NLISSRETEN"
    ws.cell(row=row_idx, column=6).value = "NLRETIR"
    for c in [3, 4, 6]:
        ws.cell(row=row_idx, column=c).font = font_bold
        ws.cell(row=row_idx, column=c).alignment = align_center
        ws.cell(row=row_idx, column=c).border = border_thin
        
    for c in [1, 2, 5, 7]:
        ws.cell(row=row_idx, column=c).border = border_thin

    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 22
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
